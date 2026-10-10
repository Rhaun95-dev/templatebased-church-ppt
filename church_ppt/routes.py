"""HTTP 라우트 (페이지 + /api/*)"""
import io
import json
import os
import tempfile
import traceback
import zipfile
from contextlib import nullcontext

from flask import Blueprint, jsonify, render_template, request, send_file
from pptx import Presentation

from .bible import load_bible_meta, lookup_verses
from .generator import generate_presentation, next_sunday_filename
from .errors import AppError, ConversionError
from .ppt_convert import convert_ppt_to_pptx
from .storage import get_storage

bp = Blueprint("main", __name__)

PPTX_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


def _asset_version():
    """프론트 파일 수정 시각 중 최댓값. 바뀌면 브라우저가 옛 JS/CSS를 쓰지 않게 한다."""
    static = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static")
    newest = 0
    for sub in ("tailwind.css", "js"):
        path = os.path.join(static, sub)
        if os.path.isdir(path):
            paths = [os.path.join(path, f) for f in os.listdir(path)]
        else:
            paths = [path]
        for p in paths:
            try:
                newest = max(newest, int(os.path.getmtime(p)))
            except OSError:
                pass
    return newest


@bp.route("/")
def index():
    return render_template("index.html", asset_v=_asset_version())


@bp.route("/api/health")
def health():
    get_storage().get_config()  # 스토리지 접근 실패 시 AppError(503) → keep-alive 겸용
    body = {"ok": True}
    # Render가 배포한 커밋 (CD 검증 워크플로가 새 버전이 떴는지 확인할 때 사용)
    commit = os.environ.get("RENDER_GIT_COMMIT")
    if commit:
        body["commit"] = commit
    return jsonify(body)


# ── Config ────────────────────────────────────


@bp.route("/api/config", methods=["GET"])
def get_config():
    st = get_storage()
    return jsonify({**st.get_config(), "mode": st.mode})


@bp.route("/api/config", methods=["POST"])
def set_config():
    get_storage().put_config(request.json or {})
    return jsonify({"ok": True})


# ── Template ──────────────────────────────────


@bp.route("/api/template", methods=["POST"])
def upload_template():
    st = get_storage()
    if st.mode != "supabase":
        return jsonify({"error": "로컬 모드에서는 설정 탭에서 템플릿 경로를 지정하세요"}), 400
    f = request.files.get("file")
    data = f.read() if f else b""
    try:
        total = len(Presentation(io.BytesIO(data)).slides)
    except Exception:
        raise ConversionError("올바른 .pptx 파일이 아니에요")
    st.put_template(data)
    return jsonify({"ok": True, "total": total})


@bp.route("/api/template-info", methods=["POST"])
def template_info():
    override = (request.json or {}).get("template_file", "")
    with get_storage().open_template(override) as template_file:
        prs = Presentation(template_file)
        slides_info = []
        for i, slide in enumerate(prs.slides):
            texts = [
                s.text_frame.text.strip()[:40]
                for s in slide.shapes
                if s.has_text_frame and s.text_frame.text.strip()
            ]
            slides_info.append(
                {
                    "index": i,
                    "number": i + 1,
                    "preview_text": (
                        " / ".join(texts[:2]) if texts else f"슬라이드 {i+1}"
                    ),
                }
            )
    return jsonify({"slides": slides_info, "total": len(slides_info)})


# ── Bible ─────────────────────────────────────


@bp.route("/api/bible-meta", methods=["GET"])
def api_bible_meta():
    return jsonify(load_bible_meta())


@bp.route("/api/bible", methods=["POST"])
def api_bible():
    data = request.json or {}
    verse_start = int(data.get("verse_start", 1))
    return jsonify(
        lookup_verses(
            data.get("book_abbr", ""),
            int(data.get("chapter", 1)),
            verse_start,
            int(data.get("verse_end", verse_start)),
        )
    )


# ── Generate ──────────────────────────────────


def _save_uploaded_hymn(upload, tmp_files):
    """업로드된 찬송가를 임시 pptx로 저장 (ppt면 변환). 저장한 경로는 tmp_files에 기록"""
    if not upload:
        return None
    ext = os.path.splitext(upload.filename.lower())[1]
    if ext not in (".ppt", ".pptx"):
        return None
    fd, path = tempfile.mkstemp(suffix=ext)
    os.close(fd)
    tmp_files.append(path)
    upload.save(path)
    # ppt면 자동 변환
    if ext == ".ppt":
        path = convert_ppt_to_pptx(path)
        tmp_files.append(path)
    return path


def _save_uploaded_template(tmp_files):
    """요청에 포함된 템플릿(사용자가 브라우저에서 고른 파일)을 임시 pptx로 저장. 없으면 None"""
    upload = request.files.get("template")
    if not upload:
        return None
    fd, path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    tmp_files.append(path)
    upload.save(path)
    if not zipfile.is_zipfile(path):
        raise ConversionError("올바른 .pptx 파일이 아니에요")
    return path


def _resolve_hymn_slots(hymn_slots, tmp_files):
    """upload_key가 있는 슬롯에 업로드 파일의 임시 경로(upload_file)를 채운다"""
    slots = []
    for slot in hymn_slots:
        # 클라이언트가 보낸 경로는 신뢰하지 않음
        slot = {k: v for k, v in slot.items() if k != "upload_file"}
        if slot.get("upload_key") and not slot.get("skip"):
            slot["upload_file"] = _save_uploaded_hymn(
                request.files.get(slot["upload_key"]), tmp_files
            )
        slots.append(slot)
    return slots


@bp.route("/api/generate", methods=["POST"])
def generate():
    upload_tmp_files = []  # 업로드된 찬송가 임시 파일 (요청 종료 시 삭제)
    try:
        if "payload" in request.form:
            data = json.loads(request.form["payload"])
        else:
            data = request.json
        # 사용자가 고른 템플릿이 있으면 그것을, 없으면 저장소의 기본 템플릿을 쓴다
        uploaded = _save_uploaded_template(upload_tmp_files)
        template_ctx = (
            nullcontext(uploaded)
            if uploaded
            else get_storage().open_template(data.get("template_file"))
        )
        with template_ctx as template_file:
            data_bytes = generate_presentation(
                template_file,
                _resolve_hymn_slots(data.get("hymn_slots", []), upload_tmp_files),
                data.get("choir", None),
                data.get("scripture", None),
                data.get("extra_verses", []),
            )

        return send_file(
            io.BytesIO(data_bytes),
            mimetype=PPTX_MIMETYPE,
            as_attachment=True,
            download_name=next_sunday_filename(),
        )

    except AppError:
        raise
    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

    finally:
        for p in upload_tmp_files:
            try:
                os.unlink(p)
            except:
                pass

from pptx import Presentation


def make_pptx(path, n_slides=3):
    prs = Presentation()
    layout = prs.slide_layouts[5]
    for i in range(n_slides):
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = f"slide {i + 1}"
    prs.save(path)
    return path

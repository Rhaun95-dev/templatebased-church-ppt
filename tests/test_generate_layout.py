import io
import os
import tempfile
import unittest
import zipfile
from collections import Counter

from pptx import Presentation

from church_ppt.generator import generate_presentation
from tests.helpers import make_pptx


def _verses(n):
    return [{"num": i, "text": f"구절{i} 본문"} for i in range(1, n + 1)]


class GenerateLayoutTests(unittest.TestCase):
    """한 번에 여러 종류의 삽입이 섞였을 때 슬라이드 수·내용·패키지 구조 확인"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.template = make_pptx(os.path.join(self.tmp, "t.pptx"), 10)
        self.hymn = make_pptx(os.path.join(self.tmp, "h.pptx"), 2)

    def _generate(self):
        choir = {
            "title_slide_index": 1,
            "lyrics_slide_index": 2,
            "song_title": "곡제목",
            "lyrics": "가사A\n\n가사B\n\n가사C",
        }
        scripture = {
            "title_slide_index": 4,
            "verses": _verses(6),
            "book_name": "요한복음",
            "chapter": 3,
            "verse_start": 1,
            "verse_end": 6,
        }
        extra = [
            {"slide_index": 7, "book_name": "시편", "chapter": 23, "verses": _verses(4)},
            {"slide_index": 7, "book_name": "롬", "chapter": 8, "verses": _verses(4)},
        ]
        slots = [{"after_slide_index": 0, "upload_file": self.hymn}]
        return generate_presentation(self.template, slots, choir, scripture, extra)

    def test_mixed_insertions(self):
        data = self._generate()

        names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        self.assertEqual([n for n, c in Counter(names).items() if c > 1], [])

        prs = Presentation(io.BytesIO(data))
        texts = [
            " ".join(s.text_frame.text for s in sl.shapes if s.has_text_frame)
            for sl in prs.slides
        ]
        # 10 + 찬송가 2 + 성가대 가사 2 + 성경 2 + 추가 구절(1 + 3)
        self.assertEqual(len(texts), 20)

        # 가사/구절이 올바른 슬라이드에 채워짐
        for lyric in ("가사A", "가사B", "가사C"):
            self.assertEqual(sum(lyric in t for t in texts), 1)
        self.assertEqual(sum("곡제목" in t for t in texts), 1)
        self.assertEqual(sum("요한복음 3장" in t for t in texts), 3)
        self.assertEqual(sum("시편 23장" in t for t in texts), 2)
        self.assertEqual(sum("롬 8장" in t for t in texts), 2)

        # 추가 구절의 템플릿 원본 슬라이드는 삭제됨
        self.assertFalse(any(t.strip() == "slide 8" for t in texts))

        # 두 번째 추가 구절 앞에 검은 슬라이드가 하나 들어감
        black = [
            sl
            for sl in prs.slides
            if sl._element.xpath('.//p:bg//a:srgbClr[@val="000000"]')
        ]
        self.assertEqual(len(black), 1)

    def test_no_insertions_returns_template_slides(self):
        data = generate_presentation(self.template, [], None, None, [])
        self.assertEqual(len(Presentation(io.BytesIO(data)).slides), 10)


if __name__ == "__main__":
    unittest.main()
